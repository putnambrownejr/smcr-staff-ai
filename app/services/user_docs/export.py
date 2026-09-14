"""Local, readable project exports; no inference or remote document service.

Word API reference: https://python-docx.readthedocs.io/en/latest/user/quickstart.html
"""

import re
from io import BytesIO
from typing import Any

from docx import Document

from app.schemas.user_docs import UserDocEntry

DRAFT_NOTICE = "DRAFT — Verify all references against current official sources before acting."
_INTERNAL = {"path", "receiptsFolder", "templateType", "kind", "archived", "editOpen", "notesOpen"}


def _label(key: str) -> str:
    return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", key).replace("_", " ").strip().capitalize()


def _sections(value: Any, level: int = 2) -> list[str]:
    lines: list[str] = []
    if isinstance(value, dict):
        for key, item in value.items():
            if key in _INTERNAL or item is None or item == "":
                continue
            lines.extend(["#" * min(level, 6) + " " + _label(str(key)), ""])
            lines.extend(_sections(item, level + 1))
    elif isinstance(value, list):
        for item in value:
            if isinstance(item, dict):
                lines.extend(_sections(item, level))
            else:
                lines.append("- " + str(item))
        lines.append("")
    else:
        lines.extend([str(value), ""])
    return lines


def render_product_markdown(entry: UserDocEntry) -> str:
    lines = ["# " + entry.title, "", "UNCLASSIFIED · Advisory draft", ""]
    if entry.body.strip():
        lines.extend([entry.body.strip(), ""])
    fields = dict(entry.fields)
    data = fields.pop("data", None)
    if data:
        lines.extend(_sections(data))
    lines.extend(_sections(fields))
    lines.extend([DRAFT_NOTICE, ""])
    return "\n".join(lines)


def render_product_docx(markdown: str) -> bytes:
    document = Document()
    document.core_properties.title = markdown.splitlines()[0].lstrip("# ")
    document.sections[0].footer.paragraphs[0].text = DRAFT_NOTICE
    for line in markdown.splitlines():
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2), level=len(heading.group(1)))
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
